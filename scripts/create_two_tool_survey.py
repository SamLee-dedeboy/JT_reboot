from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("artifacts/two_tool_interface_surveys.docx")

INK = "17252B"
ACCENT = "2D6A74"
PALE = "EAF2F3"
LINE = "AAB8BC"
MUTED = RGBColor(78, 94, 100)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=65, start=85, bottom=65, end=85):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LINE, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Arial", size=9.2, bold=False, color=INK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(paragraph, text, size=9.2, bold=False, color=INK):
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return run


def compact_paragraph(paragraph, before=0, after=0, line=1.0):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_header(doc, tool_name, description):
    p = doc.add_paragraph(style="Title")
    compact_paragraph(p, after=1)
    add_text(p, tool_name, size=18, bold=True, color="000000")
    p2 = doc.add_paragraph()
    compact_paragraph(p2, after=5, line=1.05)
    add_text(p2, description, size=8.8, color="4E5E64")


def add_section_heading(doc, text):
    p = doc.add_paragraph()
    compact_paragraph(p, before=2, after=2)
    add_text(p, text.upper(), size=8.7, bold=True, color=ACCENT)


def add_about_you(doc):
    add_section_heading(doc, "About you")
    table = doc.add_table(rows=2, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(7.35)
    set_table_borders(table, color="CBD5D8", size="5")
    questions = [
        (
            "A. What best describes your primary role?",
            "□ Researcher     □ Modeler or model developer     □ Engineer     □ Water or public agency\n"
            "□ NGO, environmental, or environmental justice organization     □ Private consultant\n"
            "□ Tribal member or representative     □ Other: ______________________________",
        ),
        (
            "B. How many years have you worked with Delta salinity, hydrodynamic modeling, or related topics?",
            "□ Less than 2     □ 2–5     □ 6–10     □ 11–20     □ More than 20",
        ),
    ]
    for row, (question, options) in zip(table.rows, questions):
        cell = row.cells[0]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_margins(cell, top=48, bottom=48, start=90, end=90)
        p = cell.paragraphs[0]
        compact_paragraph(p, after=1)
        add_text(p, question, size=8.8, bold=True)
        p2 = cell.add_paragraph()
        compact_paragraph(p2, line=1.0)
        add_text(p2, options, size=8.1)


def add_likert_intro(doc):
    p = doc.add_paragraph()
    compact_paragraph(p, before=1, after=2)
    add_text(p, "For each statement, mark one response.", size=8.2, color="4E5E64")


def add_likert_table(doc, statements):
    table = doc.add_table(rows=1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(4.95)] + [Inches(0.48)] * 5
    for col, width in zip(table.columns, widths):
        col.width = width
    headers = ["Statement", "1", "2", "3", "4", "5"]
    for idx, (cell, label) in enumerate(zip(table.rows[0].cells, headers)):
        set_cell_shading(cell, ACCENT)
        set_cell_margins(cell, top=48, bottom=48, start=55, end=55)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
        compact_paragraph(p)
        add_text(p, label, size=7.8, bold=True, color="FFFFFF")
    set_repeat_table_header(table.rows[0])
    for number, statement in enumerate(statements, start=1):
        cells = table.add_row().cells
        if number % 2 == 0:
            for cell in cells:
                set_cell_shading(cell, "F4F7F7")
        for cell in cells:
            set_cell_margins(cell, top=53, bottom=53, start=55, end=55)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cells[0].paragraphs[0]
        compact_paragraph(p, line=1.0)
        add_text(p, f"{number}. {statement}", size=8.4)
        for cell in cells[1:]:
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            compact_paragraph(p)
            add_text(p, "○", size=10.5)
    set_table_borders(table, color="C7D2D5", size="5")
    p = doc.add_paragraph()
    compact_paragraph(p, before=1, after=3)
    add_text(
        p,
        "1 Strongly disagree     2 Disagree     3 Neither agree nor disagree     4 Agree     5 Strongly agree",
        size=7.7,
        color="4E5E64",
    )


def add_open_question(doc, number, text, lines=2):
    p = doc.add_paragraph()
    compact_paragraph(p, before=2, after=1)
    add_text(p, f"{number}. {text}", size=8.8, bold=True)
    for _ in range(lines):
        line_p = doc.add_paragraph("____________________________________________________________________________________")
        compact_paragraph(line_p, after=0, line=0.85)
        for run in line_p.runs:
            set_run_font(run, size=8, color="8A999D")


def add_page_number(section, number):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    compact_paragraph(p)
    add_text(p, f"Survey {number} of 2", size=7.5, color="66777C")


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.42)
    section.bottom_margin = Inches(0.42)
    section.left_margin = Inches(0.58)
    section.right_margin = Inches(0.58)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.2)
    styles = doc.styles
    for style_name in ("Normal", "Title"):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.color.rgb = RGBColor.from_string("000000")
    styles["Normal"].font.size = Pt(9.2)


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)

    add_header(
        doc,
        "Salinity Difference Explorer Survey",
        "This survey asks how well the comparison tool supports interpretation and communication of salinity differences across scenarios, locations, and time.",
    )
    add_about_you(doc)
    add_section_heading(doc, "Your experience with the tool")
    add_likert_intro(doc)
    add_likert_table(
        doc,
        [
            "I understood what was being compared and could identify meaningful differences between scenarios.",
            "The linked map, timeline, and station distribution helped me understand relationships among scenarios, locations, and time.",
            "The tool communicated salinity differences in a way that would be useful for my work or interests.",
        ],
    )
    add_open_question(doc, 4, "What useful insight, observation, or work-related question did the tool help you explore?", lines=2)
    add_open_question(doc, 5, "If you could improve one aspect of the tool, what would you change?", lines=2)
    add_page_number(doc.sections[0], 1)

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.top_margin = Inches(0.42)
    section.bottom_margin = Inches(0.42)
    section.left_margin = Inches(0.58)
    section.right_margin = Inches(0.58)
    section.header_distance = Inches(0.2)
    section.footer_distance = Inches(0.2)
    section.footer.is_linked_to_previous = False

    add_header(
        doc,
        "Regional Salinity Pattern Explorer Survey",
        "This survey asks how clearly the interface introduces scenario strategies and communicates regional and detailed salinity patterns.",
    )
    add_about_you(doc)
    add_section_heading(doc, "Your experience with the interface")
    add_likert_intro(doc)
    add_likert_table(
        doc,
        [
            "The interface helped me understand the adaptation strategy proposed in the scenario.",
            "The regional pattern visuals made it easy to understand where and how salinity could change.",
            "The detailed pattern view helped me understand when a change occurred and how it compared with the baseline.",
            "The regional pattern visuals were visually appealing.",
        ],
    )
    add_open_question(doc, 5, "In your own words, what is the main adaptation strategy in the scenario you explored?", lines=2)
    add_open_question(
        doc,
        6,
        "What was confusing, missing, or difficult to interpret? You may comment on colors, stripes, labels, maps, timelines, or charts.",
        lines=2,
    )
    add_page_number(section, 2)

    doc.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build()
